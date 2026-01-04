#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
培训管理模块
负责培训数据管理、记录查询、分析统计等功能
"""
import os
import re
from datetime import datetime

from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, jsonify
from openpyxl import Workbook, load_workbook
import xlrd
from werkzeug.utils import secure_filename

from config.settings import (
    UPLOAD_DIR, EXPORT_DIR
)

from config.settings import APP_TITLE, EXPORT_DIR
from models.database import get_db
from .decorators import login_required, role_required
from .helpers import require_user_id, get_accessible_department_ids, validate_employee_access, build_department_filter, parse_date_filters, build_date_filter_sql, log_import_operation

# 创建 Blueprint
training_bp = Blueprint('training', __name__, url_prefix='/training')


@training_bp.route('/')
@login_required
def index():
    """培训管理主控制台"""
    from flask import session

    feature_cards = [
        {
            "title": "培训总览",
            "description": "查看培训记录情况并导出数据。",
            "endpoint": "training.records",
        },
        {
            "title": "上传Excel",
            "description": "导入月度培训Excel，自动解析培训记录并保存。",
            "endpoint": "training.upload",
        },
        {
            "title": "培训统计",
            "description": "查看培训统计分析，包含图表可视化。",
            "endpoint": "training.analytics",
        },
        {
            "title": "不合格管理",
            "description": "查看和管理培训不合格记录，跟踪整改情况。",
            "endpoint": "training.disqualified",
        },
    ]

    # 管理员可见项目管理
    if session.get('role') == 'admin':
        feature_cards.append({
            "title": "项目管理",
            "description": "管理培训项目和项目分类，规范项目命名。",
            "endpoint": "training.projects",
        })

    return render_template(
        "training_dashboard.html",
        title=f"培训管理 | {APP_TITLE}",
        feature_cards=feature_cards,
    )


@training_bp.route('/upload', methods=['GET'])
@login_required
def upload():
    """培训数据上传主页面"""
    return render_template(
        "training_upload.html",
        title=f"上传培训数据 | {APP_TITLE}",
    )


@training_bp.route('/upload/daily-report', methods=['GET', 'POST'])
@login_required
def upload_daily_report():
    """上传并导入培训日报Excel文件（支持批量.xls文件）"""
    if request.method == 'POST':
        files = request.files.getlist("files")
        if not files or all(f.filename == "" for f in files):
            flash("请选择要上传的培训日报文件。", "warning")
            return redirect(url_for("training.upload_daily_report"))

        uid = require_user_id()
        conn = get_db()
        cur = conn.cursor()

        # 获取当前用户可访问的部门ID列表（用于权限验证）
        from flask import session
        user_role = session.get('role', 'user')
        accessible_dept_ids = get_accessible_department_ids() if user_role != 'admin' else None

        # ====== 第一阶段：收集所有数据和项目名称 ======
        all_records_data = []  # 存储所有待导入的记录
        all_project_names = set()  # 存储所有项目名称
        file_errors = []

        for file_obj in files:
            if file_obj.filename == "":
                continue

            filename = secure_filename(file_obj.filename)

            # 只支持 .xls 文件
            if not filename.lower().endswith(".xls"):
                file_errors.append(f"{filename}: 仅支持 .xls 格式")
                continue

            try:
                # 读取 .xls 文件
                wb = xlrd.open_workbook(file_contents=file_obj.read(), formatting_info=False)
                sheet = wb.sheet_by_index(0)

                # 提取班组信息（第2行：填报单位）
                team_name = ""
                if sheet.nrows > 1:
                    unit_row = sheet.cell_value(1, 0)  # 第2行第1列
                    # 从"填报单位：客运二中心乘务一室2号线客车二队"中提取"2号线客车二队"
                    match = re.search(r'填报单位[：:]\s*客运二中心乘务一室(.+)', unit_row)
                    if match:
                        team_name = match.group(1).strip()

                # 提取培训日期（第3行：日期）
                training_date = None
                if sheet.nrows > 2:
                    date_row_str = str(sheet.row_values(2))  # 第3行
                    date_match = re.search(r'(\d{4})[./](\d{1,2})[./](\d{1,2})', date_row_str)
                    if date_match:
                        year, month, day = date_match.groups()
                        training_date = f"{year}-{int(month):02d}-{int(day):02d}"

                if not training_date:
                    file_errors.append(f"{filename}: 无法提取培训日期")
                    continue

                # 找到表头行（第5行，索引4）
                header_row_idx = 4
                if sheet.nrows <= header_row_idx:
                    file_errors.append(f"{filename}: 文件格式不正确")
                    continue

                # 解析表头
                header_values = sheet.row_values(header_row_idx)
                col_map = {}
                for idx, h in enumerate(header_values):
                    h_str = str(h).strip()
                    if '姓名' in h_str:
                        col_map['name'] = idx
                    elif '工号' in h_str:
                        col_map['emp_no'] = idx
                    elif '故障' in h_str:
                        col_map['project_name'] = idx  # 新格式：2025年最新
                    elif '项目类别' in h_str:
                        col_map['project_name'] = idx  # 旧格式：2025年之前
                    elif '问题类型' in h_str:
                        col_map['problem_type'] = idx
                    elif '具体问题' in h_str:
                        col_map['specific_problem'] = idx
                    elif '整改措施' in h_str:
                        col_map['corrective_measures'] = idx
                    elif '用时' in h_str:
                        col_map['time_spent'] = idx
                    elif '得分' in h_str:
                        col_map['score'] = idx
                    elif '鉴定人员' in h_str:
                        col_map['assessor'] = idx
                    elif '备注' in h_str:
                        col_map['remarks'] = idx

                if 'name' not in col_map or 'emp_no' not in col_map:
                    file_errors.append(f"{filename}: 缺少必要列（姓名、工号）")
                    continue

                # 处理数据行（从第6行开始，索引5）
                for row_idx in range(header_row_idx + 1, sheet.nrows):
                    row_values = sheet.row_values(row_idx)

                    # 跳过空行
                    if all(not str(v).strip() for v in row_values):
                        continue

                    def get_val(key):
                        idx = col_map.get(key)
                        if idx is not None and idx < len(row_values):
                            val = row_values[idx]
                            # xlrd 中数字类型需要转换
                            if isinstance(val, float) and val == int(val):
                                return int(val)
                            return val
                        return None

                    emp_no = str(get_val('emp_no') or "").strip()
                    name = str(get_val('name') or "").strip()

                    if not emp_no or not name:
                        continue

                    # 权限验证：检查该员工是否属于当前用户可访问的部门
                    if accessible_dept_ids is not None:  # 非管理员需要验证
                        cur.execute("SELECT department_id FROM employees WHERE emp_no = ?", (emp_no,))
                        emp_dept_row = cur.fetchone()

                        # 如果员工不存在或不属于可访问部门，静默跳过
                        if not emp_dept_row or emp_dept_row[0] not in accessible_dept_ids:
                            continue

                    # 提取得分
                    score_raw = get_val('score')
                    if isinstance(score_raw, (int, float)):
                        score = int(score_raw)
                    else:
                        score_match = re.search(r'\d+', str(score_raw or ""))
                        score = int(score_match.group(0)) if score_match else None

                    # 问题类型
                    problem_type = str(get_val('problem_type') or "无").strip()

                    # 判断是否合格：失格类=不合格
                    is_qualified = 0 if problem_type == "失格类" else 1
                    is_disqualified = 1 if problem_type == "失格类" else 0

                    # 备注栏判断是否补做
                    remarks = str(get_val('remarks') or "").strip()
                    is_retake = 0

                    if remarks and ("失格" in remarks or "复检" in remarks or "补做" in remarks):
                        is_retake = 1

                    # 项目名称（从"故障"列提取）
                    project_name = str(get_val('project_name') or "").strip()
                    if project_name:
                        all_project_names.add(project_name)

                    # 收集记录数据
                    record_data = {
                        'emp_no': emp_no,
                        'name': name,
                        'team_name': team_name,
                        'training_date': training_date,
                        'project_name': project_name,
                        'problem_type': problem_type,
                        'specific_problem': str(get_val('specific_problem') or ""),
                        'corrective_measures': str(get_val('corrective_measures') or ""),
                        'time_spent': str(get_val('time_spent') or ""),
                        'score': score,
                        'assessor': str(get_val('assessor') or ""),
                        'remarks': remarks,
                        'is_qualified': is_qualified,
                        'is_disqualified': is_disqualified,
                        'is_retake': is_retake,
                        'source_file': filename
                    }
                    all_records_data.append(record_data)

            except Exception as e:
                file_errors.append(f"{filename}: {str(e)}")
                continue

        # ====== 第二阶段：验证项目是否存在 ======
        if not all_records_data:
            if file_errors:
                flash(f"处理错误: {'; '.join(file_errors)}", "warning")
            else:
                flash("没有找到可导入的数据", "warning")
            return redirect(url_for("training.upload_daily_report"))

        # 查询数据库中已存在的项目
        existing_projects = {}  # {项目名称: (project_id, category_id)}
        missing_projects = []  # 不存在的项目名称列表

        for project_name in all_project_names:
            cur.execute("""
                SELECT id, category_id FROM training_projects
                WHERE name = ?
            """, (project_name,))
            row = cur.fetchone()
            if row:
                existing_projects[project_name] = (row[0], row[1])
            else:
                missing_projects.append(project_name)

        # ====== 第三阶段：如果有缺失项目，使用临时文件存储（避免session过大）======
        if missing_projects:
            import json
            import tempfile
            from pathlib import Path

            # 创建临时文件存储待导入数据
            temp_dir = Path(UPLOAD_DIR) / 'temp_imports'
            temp_dir.mkdir(exist_ok=True)

            # 生成唯一的临时文件名
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            temp_filename = f"pending_training_{uid}_{timestamp}.json"
            temp_filepath = temp_dir / temp_filename

            # 将数据序列化到临时文件
            temp_data = {
                'all_records_data': all_records_data,
                'existing_projects': existing_projects,
                'missing_projects': missing_projects,
                'file_errors': file_errors,
                'created_at': timestamp
            }

            with open(temp_filepath, 'w', encoding='utf-8') as f:
                json.dump(temp_data, f, ensure_ascii=False, indent=2)

            # session只存储临时文件路径（<100字节）
            session['pending_import_file'] = temp_filename

            # 提示用户缺失的项目信息
            project_list = "、".join(sorted(missing_projects))
            flash(f"⚠️ 发现 {len(missing_projects)} 个数据库中不存在的项目", "warning")
            flash(f"📋 缺失的项目：{project_list}", "info")
            flash(f"💡 提示：请检查项目名称是否正确，如有错误请修改Excel后重新上传", "info")
            flash(f"👉 确认无误后，请为每个项目选择分类，或将项目信息发给管理员预先创建", "warning")

            # 获取所有可用的项目分类
            cur.execute("""
                SELECT id, name FROM training_project_categories
                ORDER BY display_order ASC, name ASC
            """)
            categories = [{'id': row[0], 'name': row[1]} for row in cur.fetchall()]

            return render_template(
                'training_confirm_projects.html',
                title=f"确认项目信息 | {APP_TITLE}",
                missing_projects=sorted(missing_projects),
                categories=categories,
                record_count=len(all_records_data)
            )

        # ====== 第四阶段：直接导入（所有项目都存在） ======
        total_imported, total_skipped = _import_training_records(
            all_records_data, existing_projects, uid, conn
        )

        # 记录导入操作日志
        total_rows = len(all_records_data)
        log_import_operation(
            module='training',
            operation='import',
            file_name=f"{len(files)} files uploaded",
            total_rows=total_rows,
            success_rows=total_imported,
            failed_rows=0,
            skipped_rows=total_skipped,
            import_details={
                'imported': total_imported,
                'skipped_duplicate': total_skipped,
                'file_count': len(files),
                'file_errors': len(file_errors),
                'projects_found': len(all_project_names)
            }
        )

        # 显示结果
        if total_imported > 0:
            flash(f"成功导入 {total_imported} 条培训记录", "success")
        if total_skipped > 0:
            flash(f"跳过 {total_skipped} 条重复记录", "info")
        if file_errors:
            flash(f"处理错误: {'; '.join(file_errors)}", "warning")

        return redirect(url_for("training.records"))

    return render_template(
        "training_upload_daily.html",
        title=f"上传培训日报 | {APP_TITLE}",
    )


def _import_training_records(all_records_data, existing_projects, uid, conn):
    """
    导入培训记录的辅助函数

    Args:
        all_records_data: 所有待导入的记录数据列表
        existing_projects: 已存在的项目映射 {项目名称: (project_id, category_id)}
        uid: 当前用户ID
        conn: 数据库连接

    Returns:
        (total_imported, total_skipped): 导入成功数量和跳过数量
    """
    cur = conn.cursor()
    total_imported = 0
    total_skipped = 0

    for record in all_records_data:
        # 获取项目ID
        project_name = record['project_name']
        project_id = None

        if project_name and project_name in existing_projects:
            project_id = existing_projects[project_name][0]

        # 查找补做关联记录
        retake_of_record_id = None
        if record['is_retake']:
            date_match = re.search(r'(\d{4})[./年](\d{1,2})[./月](\d{1,2})', record['remarks'])
            if date_match:
                retake_year, retake_month, retake_day = date_match.groups()
                retake_date = f"{retake_year}-{int(retake_month):02d}-{int(retake_day):02d}"

                # 查找该人员在该日期的失格记录
                cur.execute("""
                    SELECT id FROM training_records
                    WHERE emp_no = ? AND training_date = ?
                    AND is_qualified = 0
                    LIMIT 1
                """, (record['emp_no'], retake_date))
                prev_record = cur.fetchone()
                if prev_record:
                    retake_of_record_id = prev_record[0]

        # 检查是否已存在完全相同的记录
        cur.execute("""
            SELECT COUNT(*) FROM training_records
            WHERE emp_no = ?
            AND training_date = ?
            AND project_id = ?
            AND problem_type = ?
            AND specific_problem = ?
        """, (
            record['emp_no'],
            record['training_date'],
            project_id,
            record['problem_type'],
            record['specific_problem']
        ))

        if cur.fetchone()[0] > 0:
            total_skipped += 1
            continue

        # 插入新记录
        cur.execute("""
            INSERT INTO training_records(
                emp_no, name, team_name, training_date, project_id,
                problem_type, specific_problem, corrective_measures,
                time_spent, score, assessor, remarks,
                is_qualified, is_disqualified, is_retake,
                retake_of_record_id, created_by, source_file
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            record['emp_no'],
            record['name'],
            record['team_name'],
            record['training_date'],
            project_id,
            record['problem_type'],
            record['specific_problem'],
            record['corrective_measures'],
            record['time_spent'],
            record['score'],
            record['assessor'],
            record['remarks'],
            record['is_qualified'],
            record['is_disqualified'],
            record['is_retake'],
            retake_of_record_id,
            uid,
            record['source_file']
        ))
        total_imported += 1

    conn.commit()
    return total_imported, total_skipped


@training_bp.route('/upload/confirm-projects', methods=['GET', 'POST'])
@login_required
def confirm_projects():
    """确认缺失的培训项目"""
    from flask import session
    import json
    from pathlib import Path

    # 从临时文件加载数据的辅助函数
    def load_temp_data():
        temp_filename = session.get('pending_import_file')
        if not temp_filename:
            return None

        temp_dir = Path(UPLOAD_DIR) / 'temp_imports'
        temp_filepath = temp_dir / temp_filename

        if not temp_filepath.exists():
            return None

        try:
            with open(temp_filepath, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            print(f"加载临时文件失败: {e}")
            return None

    # 清理临时文件的辅助函数
    def cleanup_temp_file():
        temp_filename = session.get('pending_import_file')
        if temp_filename:
            temp_dir = Path(UPLOAD_DIR) / 'temp_imports'
            temp_filepath = temp_dir / temp_filename
            try:
                if temp_filepath.exists():
                    temp_filepath.unlink()
            except Exception as e:
                print(f"删除临时文件失败: {e}")
            session.pop('pending_import_file', None)

    if request.method == 'POST':
        action = request.form.get('action')

        # 取消导入
        if action == 'cancel':
            cleanup_temp_file()
            flash("已取消导入", "info")
            return redirect(url_for("training.upload_daily_report"))

        # 确认并创建项目
        if action == 'confirm':
            temp_data = load_temp_data()

            if not temp_data:
                flash("会话数据已过期，请重新上传", "warning")
                cleanup_temp_file()
                return redirect(url_for("training.upload_daily_report"))

            pending_data = temp_data.get('all_records_data')
            existing_projects = temp_data.get('existing_projects', {})
            missing_projects = temp_data.get('missing_projects', [])
            file_errors = temp_data.get('file_errors', [])

            conn = get_db()
            cur = conn.cursor()
            uid = require_user_id()

            # 为每个缺失的项目创建记录
            for project_name in missing_projects:
                # 获取用户选择的分类ID
                category_id = request.form.get(f'category_{project_name}', type=int)

                if not category_id:
                    flash(f'项目"{project_name}"未选择分类', "warning")
                    return redirect(url_for("training.confirm_projects"))

                # 检查项目是否已存在（避免重复创建）
                cur.execute("""
                    SELECT id, category_id FROM training_projects
                    WHERE name = ?
                """, (project_name,))
                row = cur.fetchone()

                if row:
                    # 项目已存在，使用现有的
                    existing_projects[project_name] = (row[0], row[1])
                else:
                    # 创建新项目
                    try:
                        cur.execute("""
                            INSERT INTO training_projects (name, category_id, is_active)
                            VALUES (?, ?, 1)
                        """, (project_name, category_id))
                        new_project_id = cur.lastrowid
                        existing_projects[project_name] = (new_project_id, category_id)
                    except Exception as e:
                        conn.rollback()
                        flash(f'创建项目"{project_name}"失败: {str(e)}', "danger")
                        return redirect(url_for("training.confirm_projects"))

            # 导入所有记录
            total_imported, total_skipped = _import_training_records(
                pending_data, existing_projects, uid, conn
            )

            # 清理临时文件和session
            cleanup_temp_file()

            # 记录导入操作日志
            total_rows = len(pending_data)
            log_import_operation(
                module='training',
                operation='import_confirm_projects',
                file_name='confirmed import after project creation',
                total_rows=total_rows,
                success_rows=total_imported,
                failed_rows=0,
                skipped_rows=total_skipped,
                import_details={
                    'imported': total_imported,
                    'skipped_duplicate': total_skipped,
                    'missing_projects_created': len(missing_projects),
                    'projects_confirmed': len(existing_projects)
                }
            )

            # 显示结果
            if total_imported > 0:
                flash(f"成功导入 {total_imported} 条培训记录", "success")
            if total_skipped > 0:
                flash(f"跳过 {total_skipped} 条重复记录", "info")
            if file_errors:
                flash(f"处理错误: {'; '.join(file_errors)}", "warning")

            return redirect(url_for("training.records"))

    # GET请求：显示确认页面
    temp_data = load_temp_data()

    if not temp_data:
        flash("会话数据已过期，请重新上传", "warning")
        cleanup_temp_file()
        return redirect(url_for("training.upload_daily_report"))

    missing_projects = temp_data.get('missing_projects', [])
    record_count = len(temp_data.get('all_records_data', []))

    if not missing_projects:
        flash("没有需要确认的项目", "info")
        return redirect(url_for("training.upload_daily_report"))

    conn = get_db()
    cur = conn.cursor()

    # 获取所有可用的项目分类
    cur.execute("""
        SELECT id, name FROM training_project_categories
        ORDER BY display_order ASC, name ASC
    """)
    categories = [{'id': row[0], 'name': row[1]} for row in cur.fetchall()]

    return render_template(
        'training_confirm_projects.html',
        title=f"确认项目信息 | {APP_TITLE}",
        missing_projects=sorted(missing_projects),
        categories=categories,
        record_count=record_count
    )


@training_bp.route('/records')
@login_required
def records():
    """培训记录列表和导出"""
    from flask import session

    # 使用统一的日期筛选器
    start_date, end_date = parse_date_filters('current_month')

    name_filter = request.args.get("name", "").strip()
    qualified_filter = request.args.get("qualified")
    team_name_filter = request.args.get("team_name", "").strip()
    project_filter = request.args.get("project", "").strip()  # 项目名称筛选
    category_filter = request.args.get("category", "").strip()  # 分类筛选
    problem_type_filter = request.args.get("problem_type", "").strip()

    conn = get_db()
    cur = conn.cursor()

    # 获取当前用户角色
    user_id = session.get('user_id')
    cur.execute("SELECT role FROM users WHERE id = ?", (user_id,))
    row = cur.fetchone()
    user_role = row['role'] if row else 'user'

    # 使用新的部门过滤机制
    where_clause, join_clause, dept_params = build_department_filter('tr')

    # 构建基础查询，JOIN 项目和分类表
    base_query = f"""
        SELECT
            tr.*,
            tp.name as project_name,
            tpc.name as category_name
        FROM training_records tr
        LEFT JOIN training_projects tp ON tr.project_id = tp.id
        LEFT JOIN training_project_categories tpc ON tp.category_id = tpc.id
        {join_clause}
        WHERE {where_clause}
    """
    params = dept_params.copy()

    # 应用日期筛选
    date_conditions, date_params = build_date_filter_sql('tr.training_date', start_date, end_date)
    if date_conditions:
        base_query += " AND " + " AND ".join(date_conditions)
        params.extend(date_params)
    if name_filter:
        base_query += " AND tr.name LIKE ?"
        params.append(f"%{name_filter}%")
    if qualified_filter in ["0", "1"]:
        base_query += " AND tr.is_qualified = ?"
        params.append(int(qualified_filter))
    if team_name_filter:
        base_query += " AND tr.team_name LIKE ?"
        params.append(f"%{team_name_filter}%")
    if project_filter:
        base_query += " AND tp.name LIKE ?"
        params.append(f"%{project_filter}%")
    if category_filter:
        base_query += " AND tpc.name LIKE ?"
        params.append(f"%{category_filter}%")
    if problem_type_filter:
        base_query += " AND tr.problem_type LIKE ?"
        params.append(f"%{problem_type_filter}%")

    base_query += " ORDER BY tr.training_date DESC, tr.name"

    cur.execute(base_query, tuple(params))
    records = cur.fetchall()

    # 获取班组、项目、分类和问题类型列表用于筛选
    where_clause_for_dropdowns, _, dept_params_for_dropdowns = build_department_filter('tr')

    # 班组列表
    cur.execute(f"""
        SELECT DISTINCT tr.team_name FROM training_records tr
        {join_clause}
        WHERE {where_clause_for_dropdowns}
          AND tr.team_name IS NOT NULL AND tr.team_name != ''
        ORDER BY tr.team_name
    """, tuple(dept_params_for_dropdowns))
    team_names = [row[0] for row in cur.fetchall() if row[0]]

    # 项目名称列表（从 training_projects 获取）
    cur.execute(f"""
        SELECT DISTINCT tp.name
        FROM training_records tr
        LEFT JOIN training_projects tp ON tr.project_id = tp.id
        {join_clause}
        WHERE {where_clause_for_dropdowns}
          AND tp.name IS NOT NULL AND tp.name != ''
        ORDER BY tp.name
    """, tuple(dept_params_for_dropdowns))
    project_names = [row[0] for row in cur.fetchall() if row[0]]

    # 分类列表（从 training_project_categories 获取）
    cur.execute(f"""
        SELECT DISTINCT tpc.name
        FROM training_records tr
        LEFT JOIN training_projects tp ON tr.project_id = tp.id
        LEFT JOIN training_project_categories tpc ON tp.category_id = tpc.id
        {join_clause}
        WHERE {where_clause_for_dropdowns}
          AND tpc.name IS NOT NULL AND tpc.name != ''
        ORDER BY tpc.name
    """, tuple(dept_params_for_dropdowns))
    categories = [row[0] for row in cur.fetchall() if row[0]]

    # 问题类型列表
    cur.execute(f"""
        SELECT DISTINCT tr.problem_type FROM training_records tr
        {join_clause}
        WHERE {where_clause_for_dropdowns}
          AND tr.problem_type IS NOT NULL AND tr.problem_type != '' AND tr.problem_type != '无'
        ORDER BY tr.problem_type
    """, tuple(dept_params_for_dropdowns))
    problem_types = [row[0] for row in cur.fetchall() if row[0]]

    return render_template(
        "training_records.html",
        title=f"培训记录 | {APP_TITLE}",
        records=[dict(row) for row in records],
        start_date=start_date or "",
        end_date=end_date or "",
        name_filter=name_filter,
        qualified_filter=qualified_filter or "",
        team_name_filter=team_name_filter,
        project_filter=project_filter,
        category_filter=category_filter,
        problem_type_filter=problem_type_filter,
        team_names=team_names,
        project_names=project_names,
        categories=categories,
        problem_types=problem_types,
        user_role=user_role,
    )


@training_bp.route('/analytics')
@login_required
def analytics():
    """培训统计分析和图表"""
    return render_template(
        "training_analytics.html",
        title=f"培训统计 | {APP_TITLE}",
    )


@training_bp.route('/disqualified')
@login_required
def disqualified():
    """不合格培训记录管理"""
    # 获取筛选参数
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()
    team_filter = request.args.get("team", "").strip()
    name_filter = request.args.get("name", "").strip()
    project_filter = request.args.get("project", "").strip()
    problem_type_filter = request.args.get("problem_type", "").strip()

    conn = get_db()
    cur = conn.cursor()

    # 使用新的部门过滤机制
    where_clause, join_clause, dept_params = build_department_filter('tr')

    # 构建查询，加入 training_projects 表以获取项目名称和分类
    base_query = f"""
        SELECT tr.*, tp.name as project_name, tpc.name as category_name
        FROM training_records tr
        LEFT JOIN training_projects tp ON tr.project_id = tp.id
        LEFT JOIN training_project_categories tpc ON tp.category_id = tpc.id
        {join_clause}
        WHERE {where_clause} AND tr.is_qualified=0
    """
    params = dept_params.copy()

    # 添加筛选条件
    if start_date:
        base_query += " AND tr.training_date >= ?"
        params.append(start_date)
    if end_date:
        base_query += " AND tr.training_date <= ?"
        params.append(end_date)
    if team_filter:
        base_query += " AND tr.team_name LIKE ?"
        params.append(f"%{team_filter}%")
    if name_filter:
        base_query += " AND tr.name LIKE ?"
        params.append(f"%{name_filter}%")
    if project_filter:
        base_query += " AND tp.name LIKE ?"
        params.append(f"%{project_filter}%")
    if problem_type_filter:
        base_query += " AND tr.problem_type LIKE ?"
        params.append(f"%{problem_type_filter}%")

    base_query += " ORDER BY tr.training_date DESC"

    cur.execute(base_query, tuple(params))
    records = cur.fetchall()

    # 获取筛选选项
    cur.execute(f"""
        SELECT DISTINCT tr.team_name
        FROM training_records tr
        {join_clause}
        WHERE {where_clause} AND tr.is_qualified=0 AND tr.team_name IS NOT NULL
        ORDER BY tr.team_name
    """, dept_params)
    teams = [row[0] for row in cur.fetchall()]

    cur.execute(f"""
        SELECT DISTINCT tp.name
        FROM training_records tr
        LEFT JOIN training_projects tp ON tr.project_id = tp.id
        {join_clause}
        WHERE {where_clause} AND tr.is_qualified=0 AND tp.name IS NOT NULL
        ORDER BY tp.name
    """, dept_params)
    projects = [row[0] for row in cur.fetchall()]

    cur.execute(f"""
        SELECT DISTINCT tr.problem_type
        FROM training_records tr
        {join_clause}
        WHERE {where_clause} AND tr.is_qualified=0 AND tr.problem_type IS NOT NULL
        ORDER BY tr.problem_type
    """, dept_params)
    problem_types = [row[0] for row in cur.fetchall()]

    return render_template(
        "training_disqualified.html",
        title=f"不合格管理 | {APP_TITLE}",
        records=[dict(row) for row in records],
        teams=teams,
        projects=projects,
        problem_types=problem_types,
        start_date=start_date,
        end_date=end_date,
        team_filter=team_filter,
        name_filter=name_filter,
        project_filter=project_filter,
        problem_type_filter=problem_type_filter,
    )


@training_bp.route('/api/record/<int:record_id>')
@login_required
def get_record_detail(record_id):
    """获取培训记录详情API"""
    conn = get_db()
    cur = conn.cursor()

    # 使用部门过滤机制
    where_clause, join_clause, dept_params = build_department_filter('tr')

    query = f"""
        SELECT tr.*,
               tp.name as project_name,
               tpc.name as category_name
        FROM training_records tr
        LEFT JOIN training_projects tp ON tr.project_id = tp.id
        LEFT JOIN training_project_categories tpc ON tp.category_id = tpc.id
        {join_clause}
        WHERE {where_clause} AND tr.id = ?
    """
    params = dept_params + [record_id]

    cur.execute(query, tuple(params))
    record = cur.fetchone()

    if not record:
        return jsonify({"error": "记录不存在或无权限访问"}), 404

    return jsonify(dict(record))


@training_bp.route('/export')
@login_required
def export():
    """导出培训记录到Excel"""
    # 使用统一的日期筛选器
    start_date, end_date = parse_date_filters('current_month')

    name_filter = request.args.get("name", "").strip()
    qualified_filter = request.args.get("qualified")
    team_name_filter = request.args.get("team_name", "").strip()
    category_filter = request.args.get("category", "").strip()
    problem_type_filter = request.args.get("problem_type", "").strip()

    conn = get_db()
    cur = conn.cursor()

    # 使用新的部门过滤机制
    where_clause, join_clause, dept_params = build_department_filter('tr')

    base_query = f"""
        SELECT tr.*, tp.name as project_name, tpc.name as category_name
        FROM training_records tr
        LEFT JOIN training_projects tp ON tr.project_id = tp.id
        LEFT JOIN training_project_categories tpc ON tp.category_id = tpc.id
        {join_clause}
        WHERE {where_clause}
    """
    params = dept_params.copy()

    # 应用日期筛选
    date_conditions, date_params = build_date_filter_sql('tr.training_date', start_date, end_date)
    if date_conditions:
        base_query += " AND " + " AND ".join(date_conditions)
        params.extend(date_params)
    if name_filter:
        base_query += " AND tr.name LIKE ?"
        params.append(f"%{name_filter}%")
    if qualified_filter in ["0", "1"]:
        base_query += " AND tr.is_qualified = ?"
        params.append(int(qualified_filter))
    if team_name_filter:
        base_query += " AND tr.team_name LIKE ?"
        params.append(f"%{team_name_filter}%")
    if category_filter:
        base_query += " AND tpc.name LIKE ?"
        params.append(f"%{category_filter}%")
    if problem_type_filter:
        base_query += " AND tr.problem_type LIKE ?"
        params.append(f"%{problem_type_filter}%")

    base_query += " ORDER BY tr.training_date DESC"

    cur.execute(base_query, tuple(params))
    rows = cur.fetchall()

    if not rows:
        flash("无数据可导出", "warning")
        return redirect(url_for("training.records"))

    filename_date = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = os.path.join(EXPORT_DIR, f"培训记录_{filename_date}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "培训记录"

    headers = ["工号", "姓名", "班组", "培训日期", "项目类别", "问题类型", "具体问题", "整改措施", "用时", "得分", "鉴定人员", "备注", "是否合格"]
    ws.append(headers)

    for row in rows:
        ws.append([
            row["emp_no"], row["name"], row["team_name"] or "",
            row["training_date"], row["category_name"] or "",
            row["problem_type"] or "", row["specific_problem"] or "",
            row["corrective_measures"] or "", row["time_spent"] or "",
            row["score"] if row["score"] is not None else "",
            row["assessor"] or "", row["remarks"] or "",
            "合格" if row["is_qualified"] else "不合格"
        ])

    wb.save(xlsx_path)
    return send_file(xlsx_path, as_attachment=True, download_name=os.path.basename(xlsx_path))


@training_bp.route('/api/data')
@login_required
def api_data():
    """API端点，获取过滤后的培训数据（用于前端图表）"""
    conn = get_db()
    cur = conn.cursor()

    # 使用新的部门过滤机制
    where_clause, join_clause, dept_params = build_department_filter('tr')

    # 修改查询以包含项目和分类信息
    base_query = f"""
        SELECT
            tr.*,
            tp.name as project_name,
            tpc.name as category_name
        FROM training_records tr
        LEFT JOIN training_projects tp ON tr.project_id = tp.id
        LEFT JOIN training_project_categories tpc ON tp.category_id = tpc.id
        {join_clause}
        WHERE {where_clause}
    """
    params = dept_params.copy()

    # 使用统一的日期筛选器
    start_date, end_date = parse_date_filters('current_month')
    date_conditions, date_params = build_date_filter_sql('tr.training_date', start_date, end_date)
    if date_conditions:
        base_query += " AND " + " AND ".join(date_conditions)
        params.extend(date_params)

    name = request.args.get("name")
    if name:
        base_query += " AND tr.name LIKE ?"
        params.append(f"%{name}%")

    qualified = request.args.get("qualified")
    if qualified in ["0", "1"]:
        base_query += " AND tr.is_qualified = ?"
        params.append(int(qualified))

    base_query += " ORDER BY tr.training_date DESC"

    cur.execute(base_query, tuple(params))
    rows = cur.fetchall()

    data = [dict(row) for row in rows]
    return jsonify(data)


@training_bp.route('/records/<int:record_id>/edit', methods=['POST'])
@role_required('manager')
def edit_record(record_id):
    """编辑培训记录（仅限部门管理员及以上权限）"""
    conn = get_db()
    cur = conn.cursor()

    # 获取表单数据
    emp_no = request.form.get('emp_no', '').strip()
    name = request.form.get('name', '').strip()
    team_name = request.form.get('team_name', '').strip()
    training_date = request.form.get('training_date', '').strip()
    project_name = request.form.get('project_name', '').strip()
    problem_type = request.form.get('problem_type', '').strip()
    specific_problem = request.form.get('specific_problem', '').strip()
    corrective_measures = request.form.get('corrective_measures', '').strip()
    time_spent = request.form.get('time_spent', '').strip()
    score = request.form.get('score', '').strip()
    assessor = request.form.get('assessor', '').strip()
    remarks = request.form.get('remarks', '').strip()
    is_qualified = int(request.form.get('is_qualified', 1))

    # 验证必填字段
    if not emp_no or not name or not training_date:
        flash('工号、姓名和培训日期为必填项', 'warning')
        return redirect(url_for('training.records'))

    # 根据项目名称查找 project_id
    project_id = None
    if project_name:
        cur.execute("SELECT id FROM training_projects WHERE name = ?", (project_name,))
        project_row = cur.fetchone()
        if project_row:
            project_id = project_row['id']

    try:
        # 更新记录
        cur.execute("""
            UPDATE training_records
            SET emp_no = ?, name = ?, team_name = ?, training_date = ?,
                project_id = ?, problem_type = ?, specific_problem = ?,
                corrective_measures = ?, time_spent = ?, score = ?,
                assessor = ?, remarks = ?, is_qualified = ?
            WHERE id = ?
        """, (emp_no, name, team_name, training_date,
              project_id, problem_type, specific_problem, corrective_measures,
              time_spent, int(score) if score else None,
              assessor, remarks, is_qualified, record_id))

        conn.commit()
        flash('培训记录已更新', 'success')
    except Exception as e:
        flash(f'更新失败: {e}', 'danger')

    return redirect(url_for('training.records'))


@training_bp.route('/records/<int:record_id>/delete', methods=['POST'])
@role_required('manager')
def delete_record(record_id):
    """删除培训记录（仅限部门管理员及以上权限）"""
    conn = get_db()
    cur = conn.cursor()

    try:
        # 删除记录
        cur.execute("DELETE FROM training_records WHERE id = ?", (record_id,))
        conn.commit()
        flash('培训记录已删除', 'success')
    except Exception as e:
        flash(f'删除失败: {e}', 'danger')

    return redirect(url_for('training.records'))


@training_bp.route('/records/batch-delete', methods=['POST'])
@role_required('manager')
def batch_delete_records():
    """批量删除培训记录（仅限部门管理员及以上权限）"""
    conn = get_db()
    cur = conn.cursor()

    record_ids = request.form.getlist('record_ids')

    if not record_ids:
        flash('未选择要删除的记录', 'warning')
        return redirect(url_for('training.records'))

    try:
        # 批量删除记录
        placeholders = ','.join('?' * len(record_ids))
        cur.execute(f"DELETE FROM training_records WHERE id IN ({placeholders})", record_ids)
        conn.commit()
        flash(f'成功删除 {len(record_ids)} 条培训记录', 'success')
    except Exception as e:
        flash(f'批量删除失败: {e}', 'danger')

    return redirect(url_for('training.records'))


@training_bp.route('/test-api')
@login_required
def test_api():
    """API测试页面"""
    return render_template('test_api.html')


@training_bp.route('/debug')
@login_required
def debug_page():
    """调试页面"""
    from flask import send_file
    import os
    debug_file = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'debug_page.html')
    return send_file(debug_file)


# ==================== 项目分类管理 ====================

@training_bp.route('/project-categories')
@login_required
@role_required('admin')
def project_categories():
    """培训项目分类管理页面（仅管理员）"""
    conn = get_db()
    cur = conn.cursor()

    # 查询所有分类及其项目数量
    cur.execute("""
        SELECT
            c.id,
            c.name,
            c.description,
            c.display_order,
            c.created_at,
            COUNT(p.id) as project_count
        FROM training_project_categories c
        LEFT JOIN training_projects p ON c.id = p.category_id
        GROUP BY c.id
        ORDER BY c.display_order ASC, c.name ASC
    """)

    categories = []
    for row in cur.fetchall():
        categories.append({
            'id': row[0],
            'name': row[1],
            'description': row[2],
            'display_order': row[3],
            'created_at': row[4][:19] if row[4] else '',
            'project_count': row[5]
        })

    return render_template(
        'training_project_categories.html',
        title=f"项目分类管理 | {APP_TITLE}",
        categories=categories
    )


@training_bp.route('/project-categories/add', methods=['POST'])
@login_required
@role_required('admin')
def add_project_category():
    """添加项目分类"""
    name = request.form.get('name', '').strip()
    description = request.form.get('description', '').strip()
    display_order = request.form.get('display_order', 0, type=int)

    if not name:
        flash('分类名称不能为空', 'danger')
        return redirect(url_for('training.project_categories'))

    conn = get_db()
    cur = conn.cursor()

    try:
        cur.execute("""
            INSERT INTO training_project_categories
            (name, description, display_order)
            VALUES (?, ?, ?)
        """, (name, description, display_order))
        conn.commit()
        flash(f'项目分类"{name}"添加成功', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'添加失败: {str(e)}', 'danger')

    return redirect(url_for('training.project_categories'))


@training_bp.route('/project-categories/edit', methods=['POST'])
@login_required
@role_required('admin')
def edit_project_category():
    """编辑项目分类"""
    category_id = request.form.get('category_id', type=int)
    name = request.form.get('name', '').strip()
    description = request.form.get('description', '').strip()
    display_order = request.form.get('display_order', 0, type=int)

    if not category_id or not name:
        flash('参数错误', 'danger')
        return redirect(url_for('training.project_categories'))

    conn = get_db()
    cur = conn.cursor()

    try:
        cur.execute("""
            UPDATE training_project_categories
            SET name = ?, description = ?, display_order = ?
            WHERE id = ?
        """, (name, description, display_order, category_id))
        conn.commit()
        flash(f'项目分类"{name}"更新成功', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'更新失败: {str(e)}', 'danger')

    return redirect(url_for('training.project_categories'))


@training_bp.route('/project-categories/delete', methods=['POST'])
@login_required
@role_required('admin')
def delete_project_category():
    """删除项目分类"""
    category_id = request.form.get('category_id', type=int)

    if not category_id:
        flash('参数错误', 'danger')
        return redirect(url_for('training.project_categories'))

    conn = get_db()
    cur = conn.cursor()

    # 检查是否有关联的项目
    cur.execute("""
        SELECT COUNT(*) FROM training_projects WHERE category_id = ?
    """, (category_id,))
    project_count = cur.fetchone()[0]

    if project_count > 0:
        flash(f'该分类下有 {project_count} 个项目，无法删除', 'danger')
        return redirect(url_for('training.project_categories'))

    try:
        cur.execute("DELETE FROM training_project_categories WHERE id = ?", (category_id,))
        conn.commit()
        flash('项目分类删除成功', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'删除失败: {str(e)}', 'danger')

    return redirect(url_for('training.project_categories'))


# ==================== 项目管理 ====================

@training_bp.route('/projects')
@login_required
@role_required('admin')
def projects():
    """培训项目管理页面（仅管理员）"""
    conn = get_db()
    cur = conn.cursor()

    # 获取筛选条件
    category_id = request.args.get('category_id', type=int)
    is_active = request.args.get('is_active')
    search = request.args.get('search', '').strip()

    # 查询所有分类
    cur.execute("SELECT id, name FROM training_project_categories ORDER BY display_order ASC, name ASC")
    categories = [{'id': row[0], 'name': row[1]} for row in cur.fetchall()]

    # 构建查询
    query = """
        SELECT
            p.id,
            p.name,
            p.category_id,
            c.name as category_name,
            p.description,
            p.is_active,
            p.created_at,
            COUNT(tr.id) as record_count
        FROM training_projects p
        LEFT JOIN training_project_categories c ON p.category_id = c.id
        LEFT JOIN training_records tr ON p.id = tr.project_id
        WHERE 1=1
    """
    params = []

    if category_id:
        query += " AND p.category_id = ?"
        params.append(category_id)

    if is_active is not None and is_active != '':
        query += " AND p.is_active = ?"
        params.append(int(is_active))

    if search:
        query += " AND p.name LIKE ?"
        params.append(f'%{search}%')

    query += " GROUP BY p.id ORDER BY c.display_order ASC, p.name ASC"

    cur.execute(query, params)

    projects = []
    for row in cur.fetchall():
        projects.append({
            'id': row[0],
            'name': row[1],
            'category_id': row[2],
            'category_name': row[3] or '未分类',
            'description': row[4],
            'is_active': row[5],
            'created_at': row[6],
            'record_count': row[7]
        })

    return render_template(
        'training_projects.html',
        title=f"项目管理 | {APP_TITLE}",
        projects=projects,
        categories=categories
    )


@training_bp.route('/projects/add', methods=['POST'])
@login_required
@role_required('admin')
def add_project():
    """添加培训项目"""
    name = request.form.get('name', '').strip()
    category_id = request.form.get('category_id', type=int)
    description = request.form.get('description', '').strip()
    is_active = 1 if request.form.get('is_active') else 0

    if not name or not category_id:
        flash('项目名称和分类不能为空', 'danger')
        return redirect(url_for('training.projects'))

    conn = get_db()
    cur = conn.cursor()

    try:
        cur.execute("""
            INSERT INTO training_projects
            (name, category_id, description, is_active)
            VALUES (?, ?, ?, ?)
        """, (name, category_id, description, is_active))
        conn.commit()
        flash(f'培训项目"{name}"添加成功', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'添加失败: {str(e)}', 'danger')

    return redirect(url_for('training.projects'))


@training_bp.route('/projects/edit', methods=['POST'])
@login_required
@role_required('admin')
def edit_project():
    """编辑培训项目"""
    project_id = request.form.get('project_id', type=int)
    name = request.form.get('name', '').strip()
    category_id = request.form.get('category_id', type=int)
    description = request.form.get('description', '').strip()
    is_active = 1 if request.form.get('is_active') else 0

    if not project_id or not name or not category_id:
        flash('参数错误', 'danger')
        return redirect(url_for('training.projects'))

    conn = get_db()
    cur = conn.cursor()

    try:
        cur.execute("""
            UPDATE training_projects
            SET name = ?, category_id = ?, description = ?, is_active = ?
            WHERE id = ?
        """, (name, category_id, description, is_active, project_id))
        conn.commit()
        flash(f'培训项目"{name}"更新成功', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'更新失败: {str(e)}', 'danger')

    return redirect(url_for('training.projects'))


@training_bp.route('/projects/delete', methods=['POST'])
@login_required
@role_required('admin')
def delete_project():
    """删除培训项目"""
    project_id = request.form.get('project_id', type=int)

    if not project_id:
        flash('参数错误', 'danger')
        return redirect(url_for('training.projects'))

    conn = get_db()
    cur = conn.cursor()

    # 检查是否有关联的培训记录
    cur.execute("SELECT COUNT(*) FROM training_records WHERE project_id = ?", (project_id,))
    record_count = cur.fetchone()[0]

    if record_count > 0:
        flash(f'该项目有 {record_count} 条培训记录，无法删除', 'danger')
        return redirect(url_for('training.projects'))

    try:
        cur.execute("DELETE FROM training_projects WHERE id = ?", (project_id,))
        conn.commit()
        flash('培训项目删除成功', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'删除失败: {str(e)}', 'danger')

    return redirect(url_for('training.projects'))


@training_bp.route('/projects/batch-delete', methods=['POST'])
@login_required
@role_required('admin')
def batch_delete_projects():
    """批量删除培训项目"""
    project_ids = request.form.getlist('project_ids')

    if not project_ids:
        flash('未选择要删除的项目', 'warning')
        return redirect(url_for('training.projects'))

    conn = get_db()
    cur = conn.cursor()

    deleted_count = 0
    skipped_count = 0
    errors = []

    for project_id in project_ids:
        try:
            # 检查是否有关联的培训记录
            cur.execute("SELECT COUNT(*) FROM training_records WHERE project_id = ?", (project_id,))
            record_count = cur.fetchone()[0]

            if record_count > 0:
                # 获取项目名称用于提示
                cur.execute("SELECT name FROM training_projects WHERE id = ?", (project_id,))
                row = cur.fetchone()
                if row:
                    errors.append(f'"{row[0]}"有{record_count}条记录')
                skipped_count += 1
                continue

            # 删除项目
            cur.execute("DELETE FROM training_projects WHERE id = ?", (project_id,))
            deleted_count += 1

        except Exception as e:
            errors.append(f'ID {project_id}: {str(e)}')
            skipped_count += 1

    conn.commit()

    # 显示结果
    if deleted_count > 0:
        flash(f'成功删除 {deleted_count} 个项目', 'success')
    if skipped_count > 0:
        flash(f'跳过 {skipped_count} 个项目（{"; ".join(errors[:5])}）', 'warning')

    return redirect(url_for('training.projects'))


@training_bp.route('/projects/batch-add', methods=['POST'])
@login_required
@role_required('admin')
def batch_add_projects():
    """批量添加培训项目"""
    batch_data = request.form.get('batch_data', '').strip()
    default_category_id = request.form.get('default_category_id', type=int)
    is_active = 1 if request.form.get('is_active') else 0

    if not batch_data:
        flash('请粘贴要添加的数据', 'warning')
        return redirect(url_for('training.projects'))

    conn = get_db()
    cur = conn.cursor()

    # 获取现有分类映射 {分类名称: 分类ID}
    cur.execute("SELECT id, name FROM training_project_categories")
    category_map = {row[1]: row[0] for row in cur.fetchall()}

    # 解析数据
    lines = batch_data.split('\n')
    added_count = 0
    skipped_count = 0
    new_categories = []
    errors = []

    for line_no, line in enumerate(lines, 1):
        line = line.strip()
        if not line:
            continue

        # 解析每行数据
        parts = [p.strip() for p in line.split('\t')]

        category_name = None
        project_name = None

        if len(parts) >= 2:
            # 两列数据：分类 + 项目名称
            category_name = parts[0] if parts[0] else None
            project_name = parts[1]
        elif len(parts) == 1:
            # 只有一列：项目名称
            project_name = parts[0]
        else:
            errors.append(f'第{line_no}行：格式错误')
            skipped_count += 1
            continue

        if not project_name:
            errors.append(f'第{line_no}行：项目名称为空')
            skipped_count += 1
            continue

        # 确定分类ID
        category_id = None

        if category_name:
            # 检查分类是否存在
            if category_name in category_map:
                category_id = category_map[category_name]
            else:
                # 创建新分类
                try:
                    # 获取当前最大display_order
                    cur.execute("SELECT COALESCE(MAX(display_order), 0) FROM training_project_categories")
                    max_order = cur.fetchone()[0]

                    cur.execute("""
                        INSERT INTO training_project_categories (name, display_order)
                        VALUES (?, ?)
                    """, (category_name, max_order + 1))
                    category_id = cur.lastrowid
                    category_map[category_name] = category_id
                    new_categories.append(category_name)
                except Exception as e:
                    errors.append(f'第{line_no}行：创建分类"{category_name}"失败 - {str(e)}')
                    skipped_count += 1
                    continue
        elif default_category_id:
            # 使用默认分类
            category_id = default_category_id
        else:
            # 没有分类且没有默认分类，跳过
            errors.append(f'第{line_no}行：未指定分类')
            skipped_count += 1
            continue

        # 添加项目
        try:
            cur.execute("""
                INSERT INTO training_projects (name, category_id, is_active)
                VALUES (?, ?, ?)
            """, (project_name, category_id, is_active))
            added_count += 1
        except Exception as e:
            errors.append(f'第{line_no}行：添加项目"{project_name}"失败 - {str(e)}')
            skipped_count += 1

    conn.commit()

    # 显示结果
    if added_count > 0:
        flash(f'成功添加 {added_count} 个项目', 'success')
    if new_categories:
        flash(f'自动创建了 {len(new_categories)} 个新分类：{", ".join(new_categories)}', 'info')
    if skipped_count > 0:
        error_msg = '; '.join(errors[:5])
        if len(errors) > 5:
            error_msg += f' 等共{len(errors)}个错误'
        flash(f'跳过 {skipped_count} 条数据（{error_msg}）', 'warning')

    return redirect(url_for('training.projects'))
